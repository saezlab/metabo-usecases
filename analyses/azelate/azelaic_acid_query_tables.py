from __future__ import annotations

import json
import re
import time
from collections import OrderedDict
from dataclasses import dataclass
from datetime import datetime, timezone
from itertools import zip_longest
from pathlib import Path
from typing import Any

import pandas as pd
import requests


API_BASE = "https://dev.omnipathdb.org/api"

QUERY_TERMS = [
    "Azelaic acid (PTF52203)",
    "Azelaic acid; nonanedioic acid",
    "Azelate",
]

FOCAL_INCHIKEY = "BDJRBEYXGGNYIS-UHFFFAOYSA-N"
ANION_INCHIKEY = "BDJRBEYXGGNYIS-UHFFFAOYSA-L"
SEED_SELECTION_RULES = (
    {
        "seed_key": "ptfi_azelaic_acid_ptf52203",
        "seed_display_label": "Azelaic acid (PTF52203)",
        "seed_resolution_status": "unresolved",
        "required_source": "ptfi",
        "required_identifier": "PTF52203",
    },
    {
        "seed_key": "recon3d_azelaic_acid_nonanedioic_acid",
        "seed_display_label": "Azelaic acid; nonanedioic acid",
        "seed_resolution_status": "unresolved",
        "required_source": "recon3d",
        "required_identifier": "Azelaic acid; nonanedioic acid",
    },
    {
        "seed_key": "resolved_azelate_standard_inchikey",
        "seed_display_label": "Azelate",
        "seed_resolution_status": "resolved",
        "required_canonical": FOCAL_INCHIKEY,
        "required_identifier": "Azelate",
    },
)
REJECTED_SEED_TEXT_TOKENS = (
    "3-methylazelaic",
    "methylazelaic",
)

MACDB_TERM_MAP = OrderedDict(
    [
        ("Type:OM:1252", "disease_type"),
        ("Subtype:OM:1253", "disease_subtype"),
        ("Tissue:OM:0764", "tissue_original"),
        ("Experimental Method:OM:0687", "experimental_method"),
        ("Case Sample Count:OM:0773", "case_sample_count"),
        ("Control Sample Count:OM:0774", "control_sample_count"),
        ("Pubmed:MI:0446", "pubmed_ids"),
        ("Pubmed Central:MI:1042", "pmc_ids"),
        ("Description:OM:0766", "study_description"),
        ("Case Description:OM:0775", "case_description"),
        ("Control Description:OM:0776", "control_description"),
        ("Contrast P Val:OM:0767", "contrast_p_value"),
        ("Contrast Logfc:OM:0768", "contrast_logfc"),
        ("Conclusion:OM:0778", "conclusion"),
        ("Metac:OM:0135", "metac_id"),
        ("Phenotype Result:MI:2283", "phenotype_result"),
    ]
)

METABOLIC_SOURCES = {"metatlas", "rhea", "kegg", "recon3d", "reactome"}


@dataclass
class PipelineResult:
    tables: dict[str, pd.DataFrame]
    workbook_path: Path
    csv_dir: Path
    figure_paths: dict[str, Path]
    manifest_path: Path
    output_dir: Path


class OmniPathFastApiClient:
    def __init__(self, base_url: str = API_BASE, timeout: int = 45, retries: int = 2, pause: float = 0.0):
        self.base_url = base_url.rstrip("/")
        self.timeout = timeout
        self.retries = retries
        self.pause = pause
        self.session = requests.Session()
        self.request_log: list[dict[str, Any]] = []

    def _url(self, path: str) -> str:
        if path.startswith("http://") or path.startswith("https://"):
            return path
        return f"{self.base_url}/{path.lstrip('/')}"

    def _request(self, method: str, path: str, **kwargs: Any) -> Any:
        url = self._url(path)
        last_error: Exception | None = None
        for attempt in range(1, self.retries + 1):
            started = time.time()
            try:
                response = self.session.request(method, url, timeout=self.timeout, **kwargs)
                elapsed = time.time() - started
                self.request_log.append(
                    {
                        "method": method.upper(),
                        "url": response.url,
                        "path": path,
                        "status_code": response.status_code,
                        "elapsed_seconds": round(elapsed, 3),
                        "attempt": attempt,
                    }
                )
                response.raise_for_status()
                if self.pause:
                    time.sleep(self.pause)
                return response.json()
            except Exception as exc:
                last_error = exc
                if attempt < self.retries:
                    time.sleep(min(2**attempt, 10))
        raise RuntimeError(f"{method.upper()} {url} failed after {self.retries} attempts: {last_error}") from last_error

    def get(self, path: str, params: dict[str, Any] | None = None) -> Any:
        return self._request("GET", path, params=params)

    def post(self, path: str, payload: dict[str, Any] | None = None) -> Any:
        return self._request("POST", path, json=payload or {})


def uniq_pipe(values: Any) -> str:
    if values is None:
        return ""
    if isinstance(values, str):
        return values
    seen: OrderedDict[str, None] = OrderedDict()
    for value in values:
        if pd.isna(value):
            continue
        text = str(value).strip()
        if text and text.lower() != "nan":
            seen[text] = None
    return " | ".join(seen.keys())


def split_pipe(value: Any) -> list[str]:
    if pd.isna(value):
        return []
    parts = [part.strip() for part in str(value).split("|")]
    return [part for part in parts if part and part.lower() != "nan"]


def source_has(source_pipe: Any, source: str) -> bool:
    return source in {part.strip() for part in str(source_pipe or "").split("|") if part.strip()}


def clean_identifier_type(value: Any) -> str:
    text = str(value or "").split(":")[0].replace("_", " ").strip()
    replacements = {
        "Cas": "CAS",
        "Chembl": "ChEMBL",
        "Chembl Compound": "ChEMBL Compound",
        "Chembl Target": "ChEMBL Target",
        "Chebi": "ChEBI",
        "Drugbank": "DrugBank",
        "Foodb": "FooDB",
        "Hmdb": "HMDB",
        "Human Gem Metabolite": "Human-GEM Metabolite",
        "Inchi": "InChI",
        "Iupac": "IUPAC",
        "Iupac Name": "IUPAC Name",
        "Iupac Traditional Name": "IUPAC Traditional Name",
        "Kegg Compound": "KEGG Compound",
        "Lipidmaps": "LIPID MAPS",
        "Molecular Formula": "Molecular Formula",
        "Pubchem Compound": "PubChem Compound",
        "Pubchem": "PubChem",
        "Refmet": "RefMet",
        "Smiles": "SMILES",
        "Standard Inchi Key": "Standard InChI Key",
        "Swisslipids": "SwissLipids",
    }
    return replacements.get(text.title(), replacements.get(text, text))


def identifier_category(identifier_type_clean: str) -> str:
    text = identifier_type_clean.lower()
    if any(token in text for token in ["smiles", "inchi", "iupac", "formula"]):
        return "Structure"
    return "ID"


def safe_json(value: Any) -> str:
    if value in (None, ""):
        return ""
    return json.dumps(value, ensure_ascii=False, sort_keys=True)


def preferred_identifier(entity: dict[str, Any], prefixes: tuple[str, ...]) -> str | None:
    identifiers = entity.get("identifiers") or []
    for prefix in prefixes:
        for item in identifiers:
            if str(item.get("identifierType", "")).startswith(prefix):
                value = item.get("identifier")
                if value:
                    return str(value)
    return None


def entity_label(entity: dict[str, Any] | None) -> str:
    if not entity:
        return ""
    label = preferred_identifier(
        entity,
        (
            "Name:OM:0200",
            "Name:OM:0202",
            "Gene Name Primary",
            "Iupac Traditional Name",
            "Iupac Name",
            "Synonym",
        ),
    )
    return label or str(entity.get("canonicalIdentifier") or entity.get("entityPk") or "")


def entity_summary(entity: dict[str, Any]) -> dict[str, Any]:
    return {
        "entity_pk": entity.get("entityPk"),
        "label": entity_label(entity),
        "canonical_identifier": entity.get("canonicalIdentifier"),
        "canonical_identifier_type": entity.get("canonicalIdentifierType"),
        "entity_type": entity.get("entityType"),
        "taxonomy_id": entity.get("taxonomyId"),
        "sources": uniq_pipe(entity.get("sources") or []),
        "relation_count": entity.get("relationCount"),
        "n_identifiers": len(entity.get("identifiers") or []),
    }


def flatten_identifiers(entities: list[dict[str, Any]], entity_scope: str) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    for entity in entities:
        base = entity_summary(entity)
        for item in entity.get("identifiers") or []:
            identifier_type = item.get("identifierType")
            rows.append(
                {
                    "entity_scope": entity_scope,
                    **base,
                    "identifier": item.get("identifier"),
                    "identifier_type": identifier_type,
                    "identifier_type_clean": clean_identifier_type(identifier_type),
                    "identifier_record_id": item.get("id"),
                }
            )
    return pd.DataFrame(rows)


def explode_entity_sources(entities: list[dict[str, Any]], entity_scope: str) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    for entity in entities:
        base = entity_summary(entity)
        for source in entity.get("sources") or []:
            rows.append({"entity_scope": entity_scope, **base, "source": source})
    return pd.DataFrame(rows)


def paged_entities_search(
    client: OmniPathFastApiClient,
    query: str,
    limit: int = 100,
    max_records: int = 100,
) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    offset = 0
    while True:
        page_limit = min(limit, max_records - len(out))
        if page_limit <= 0:
            break
        payload = client.get("/entities/search", params={"q": query, "limit": page_limit, "offset": offset})
        rows = payload.get("entities") or []
        out.extend(rows)
        total = int(payload.get("total") or len(out))
        if not rows or len(out) >= total or len(out) >= max_records:
            break
        offset += page_limit
    return out


def search_entities(client: OmniPathFastApiClient, query_terms: list[str]) -> tuple[pd.DataFrame, dict[str, dict[str, Any]]]:
    rows: list[dict[str, Any]] = []
    entity_map: dict[str, dict[str, Any]] = {}
    for query in query_terms:
        entities = paged_entities_search(client, query)
        for rank, entity in enumerate(entities, start=1):
            entity_map[entity["entityPk"]] = entity
            rows.append(
                {
                    "query_term": query,
                    "search_rank": rank,
                    "search_total_for_query": len(entities),
                    **entity_summary(entity),
                }
            )
    df = pd.DataFrame(rows)
    if not df.empty:
        df = df.drop_duplicates(["query_term", "entity_pk"]).sort_values(["query_term", "search_rank"])
    return df.reset_index(drop=True), entity_map


def resolve_entities(client: OmniPathFastApiClient, query_terms: list[str]) -> tuple[pd.DataFrame, pd.DataFrame, dict[str, dict[str, Any]]]:
    payload = client.post("/entities/resolve", {"identifiers": query_terms, "limit": 100})
    match_rows: list[dict[str, Any]] = []
    candidate_rows: list[dict[str, Any]] = []
    entity_map: dict[str, dict[str, Any]] = {}

    for entity in payload.get("entities") or []:
        entity_map[entity["entityPk"]] = entity

    for match in payload.get("matches") or []:
        match_rows.append(
            {
                "query_term": match.get("identifier"),
                "entity_pks": uniq_pipe(match.get("entityPks") or []),
                "best_entity_pk": match.get("bestEntityPk"),
                "ambiguous": match.get("ambiguous"),
                "n_candidates": len(match.get("candidates") or []),
            }
        )
        for rank, candidate in enumerate(match.get("candidates") or [], start=1):
            if candidate.get("entityPk"):
                entity_map[candidate["entityPk"]] = candidate
            candidate_rows.append(
                {
                    "query_term": match.get("identifier"),
                    "candidate_rank": rank,
                    "best_entity_pk": match.get("bestEntityPk"),
                    "ambiguous": match.get("ambiguous"),
                    "match_identifier": candidate.get("matchIdentifier"),
                    "match_identifier_type": candidate.get("matchIdentifierType"),
                    "match_kind": candidate.get("matchKind"),
                    "score": candidate.get("score"),
                    **entity_summary(candidate),
                }
            )
    return pd.DataFrame(match_rows), pd.DataFrame(candidate_rows), entity_map


def select_seed_entities(
    search_hits: pd.DataFrame,
    search_entities_by_pk: dict[str, dict[str, Any]],
    resolve_entities_by_pk: dict[str, dict[str, Any]],
) -> dict[str, dict[str, Any]]:
    selected: OrderedDict[str, dict[str, Any]] = OrderedDict()

    if not search_hits.empty:
        for _, row in search_hits.iterrows():
            entity = search_entities_by_pk.get(row["entity_pk"]) or resolve_entities_by_pk.get(row["entity_pk"])
            if entity and seed_selection_match(entity):
                selected[row["entity_pk"]] = entity

    for entity_map in (search_entities_by_pk, resolve_entities_by_pk):
        for pk, entity in entity_map.items():
            if seed_selection_match(entity):
                selected[pk] = entity

    validate_seed_selection(selected)
    return selected


def entity_identifier_values(entity: dict[str, Any]) -> list[str]:
    values: list[str] = []
    if entity.get("canonicalIdentifier"):
        values.append(str(entity.get("canonicalIdentifier")))
    values.extend(str(item.get("identifier") or "") for item in entity.get("identifiers") or [])
    return [value for value in values if value]


def entity_identifier_set(entity: dict[str, Any]) -> set[str]:
    return {value.casefold() for value in entity_identifier_values(entity)}


def has_entity_identifier(entity: dict[str, Any], identifier: str) -> bool:
    wanted = identifier.casefold()
    return any(value.casefold() == wanted for value in entity_identifier_values(entity))


def seed_selection_match(entity: dict[str, Any]) -> dict[str, str] | None:
    if str(entity.get("entityType") or "") != "Chemical:OM:0037":
        return None

    text = " | ".join([entity_label(entity), *entity_identifier_values(entity)]).casefold()
    if any(token in text for token in REJECTED_SEED_TEXT_TOKENS):
        return None

    sources = {str(source).casefold() for source in entity.get("sources") or []}
    canonical = str(entity.get("canonicalIdentifier") or "")
    for seed_order, rule in enumerate(SEED_SELECTION_RULES, start=1):
        required_source = rule.get("required_source")
        if required_source and str(required_source).casefold() not in sources:
            continue
        required_canonical = rule.get("required_canonical")
        if required_canonical and canonical != required_canonical:
            continue
        required_identifier = rule.get("required_identifier")
        if required_identifier and not has_entity_identifier(entity, str(required_identifier)):
            continue
        return {
            "seed_key": str(rule["seed_key"]),
            "seed_display_label": str(rule["seed_display_label"]),
            "seed_resolution_status": str(rule["seed_resolution_status"]),
            "seed_order": str(seed_order),
        }
    return None


def validate_seed_selection(selected: dict[str, dict[str, Any]]) -> None:
    matches: dict[str, list[str]] = {str(rule["seed_key"]): [] for rule in SEED_SELECTION_RULES}
    rejected: list[str] = []
    for pk, entity in selected.items():
        match = seed_selection_match(entity)
        if not match:
            rejected.append(pk)
            continue
        matches[match["seed_key"]].append(pk)

    missing = [rule["seed_display_label"] for rule in SEED_SELECTION_RULES if not matches[str(rule["seed_key"])]]
    duplicated = {key: pks for key, pks in matches.items() if len(pks) > 1}
    if missing or duplicated or rejected or len(selected) != len(SEED_SELECTION_RULES):
        selected_rows = []
        for pk, entity in selected.items():
            match = seed_selection_match(entity) or {}
            selected_rows.append(
                {
                    "entity_pk": pk,
                    "label": entity_label(entity),
                    "canonical_identifier": entity.get("canonicalIdentifier"),
                    "sources": uniq_pipe(entity.get("sources") or []),
                    "seed_key": match.get("seed_key", ""),
                    "seed_display_label": match.get("seed_display_label", ""),
                    "seed_resolution_status": match.get("seed_resolution_status", ""),
                }
            )
        diagnostics = pd.DataFrame(selected_rows).to_string(index=False) if selected_rows else "<none>"
        raise RuntimeError(
            "Seed selection did not resolve exactly the three required azelaic-acid entities. "
            f"Missing: {missing or 'none'}; duplicated: {duplicated or 'none'}; "
            f"rejected: {rejected or 'none'}; selected_count={len(selected)}.\n{diagnostics}"
        )


def seed_selection_metadata(entity: dict[str, Any]) -> dict[str, str]:
    match = seed_selection_match(entity)
    if not match:
        return {"seed_key": "", "seed_display_label": "", "seed_resolution_status": "", "seed_order": ""}
    return match


def summarize_resolution_status(statuses: list[str]) -> str:
    unique_statuses = [status for status in OrderedDict((status, None) for status in statuses if status).keys()]
    if not unique_statuses:
        return ""
    if len(unique_statuses) == 1:
        return unique_statuses[0]
    return "mixed"


def add_query_seed_metadata(
    relation_query_hits: pd.DataFrame,
    seed_metadata_by_pk: dict[str, dict[str, str]],
) -> pd.DataFrame:
    if relation_query_hits.empty:
        return relation_query_hits
    out = relation_query_hits.copy()
    out["query_seed_display_label"] = out["query_entity_pk"].map(
        lambda pk: seed_metadata_by_pk.get(str(pk), {}).get("seed_display_label", "")
    )
    out["query_seed_resolution_status"] = out["query_entity_pk"].map(
        lambda pk: seed_metadata_by_pk.get(str(pk), {}).get("seed_resolution_status", "")
    )
    return out


def add_seed_metadata_columns(
    df: pd.DataFrame,
    seed_metadata_by_pk: dict[str, dict[str, str]],
    key_col: str = "entity_pk",
    prefix: str = "seed_",
) -> pd.DataFrame:
    if df.empty or key_col not in df.columns:
        return df
    out = df.copy()
    out[f"{prefix}key"] = out[key_col].map(lambda pk: seed_metadata_by_pk.get(str(pk), {}).get("seed_key", ""))
    out[f"{prefix}display_label"] = out[key_col].map(
        lambda pk: seed_metadata_by_pk.get(str(pk), {}).get("seed_display_label", "")
    )
    out[f"{prefix}resolution_status"] = out[key_col].map(
        lambda pk: seed_metadata_by_pk.get(str(pk), {}).get("seed_resolution_status", "")
    )
    out[f"{prefix}order"] = out[key_col].map(lambda pk: seed_metadata_by_pk.get(str(pk), {}).get("seed_order", ""))
    return out


def add_relation_seed_metadata(df: pd.DataFrame, relation_table: pd.DataFrame) -> pd.DataFrame:
    if df.empty or relation_table.empty or "relation_pk" not in df.columns:
        return df
    cols = [
        "relation_pk",
        "matched_seed_entity_pks",
        "matched_seed_display_labels",
        "matched_seed_resolution_status",
        "matched_seed_resolution_statuses",
    ]
    available_cols = [col for col in cols if col in relation_table.columns]
    metadata = relation_table[available_cols].drop_duplicates("relation_pk")
    return df.merge(metadata, on="relation_pk", how="left")


def hydrate_entities(client: OmniPathFastApiClient, entity_pks: list[str], chunk_size: int = 100) -> dict[str, dict[str, Any]]:
    hydrated: dict[str, dict[str, Any]] = {}
    unique_pks = list(OrderedDict((pk, None) for pk in entity_pks if pk).keys())
    for start in range(0, len(unique_pks), chunk_size):
        chunk = unique_pks[start : start + chunk_size]
        payload = client.post("/entities/by-pks", {"entityPks": chunk})
        for entity in payload.get("entities") or []:
            hydrated[entity["entityPk"]] = entity
    return hydrated


def fetch_relations_for_entities(client: OmniPathFastApiClient, entity_pks: list[str], limit: int = 500) -> tuple[pd.DataFrame, dict[str, dict[str, Any]]]:
    hit_rows: list[dict[str, Any]] = []
    relation_map: dict[str, dict[str, Any]] = {}
    for entity_pk in entity_pks:
        offset = 0
        while True:
            payload = client.get(
                "/relations/search",
                params={"entityPks": entity_pk, "limit": limit, "offset": offset},
            )
            rows = payload.get("relations") or []
            total = int(payload.get("total") or len(rows))
            for relation in rows:
                relation_map[relation["relationPk"]] = relation
                hit_rows.append(
                    {
                        "query_entity_pk": entity_pk,
                        "relation_pk": relation.get("relationPk"),
                        "predicate": relation.get("predicate"),
                        "relation_category": relation.get("relationCategory"),
                        "sources": uniq_pipe(relation.get("sources") or []),
                        "evidence_count": relation.get("evidenceCount"),
                    }
                )
            if not rows or offset + limit >= total:
                break
            offset += limit
    return pd.DataFrame(hit_rows).drop_duplicates().reset_index(drop=True), relation_map


def relation_participant_pks(relation_map: dict[str, dict[str, Any]]) -> list[str]:
    pks: list[str] = []
    for relation in relation_map.values():
        for key in ("subjectEntityPk", "objectEntityPk"):
            if relation.get(key):
                pks.append(relation[key])
    return list(OrderedDict((pk, None) for pk in pks).keys())


def build_direct_relations(
    relation_map: dict[str, dict[str, Any]],
    seed_entity_pks: set[str],
    seed_metadata_by_pk: dict[str, dict[str, str]],
    entity_map: dict[str, dict[str, Any]],
) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    for relation in relation_map.values():
        subject_pk = relation.get("subjectEntityPk")
        object_pk = relation.get("objectEntityPk")
        subject = entity_map.get(subject_pk) or relation.get("subjectEntity") or {}
        obj = entity_map.get(object_pk) or relation.get("objectEntity") or {}
        matched_seed_pks = [pk for pk in (subject_pk, object_pk) if pk in seed_entity_pks]
        matched_seed_metadata = [seed_metadata_by_pk.get(pk, {}) for pk in matched_seed_pks]
        matched_seed_statuses = [metadata.get("seed_resolution_status", "") for metadata in matched_seed_metadata]
        other_pks = [pk for pk in (subject_pk, object_pk) if pk and pk not in seed_entity_pks]
        other_entities = [entity_map.get(pk) for pk in other_pks]
        rows.append(
            {
                "relation_pk": relation.get("relationPk"),
                "predicate": relation.get("predicate"),
                "relation_category": relation.get("relationCategory"),
                "sources": uniq_pipe(relation.get("sources") or []),
                "evidence_count": relation.get("evidenceCount"),
                "participant_types": uniq_pipe(relation.get("participantTypes") or []),
                "subject_entity_pk": subject_pk,
                "subject_label": entity_label(subject),
                "subject_canonical": subject.get("canonicalIdentifier"),
                "subject_entity_type": subject.get("entityType"),
                "subject_taxonomy_id": subject.get("taxonomyId"),
                "object_entity_pk": object_pk,
                "object_label": entity_label(obj),
                "object_canonical": obj.get("canonicalIdentifier"),
                "object_entity_type": obj.get("entityType"),
                "object_taxonomy_id": obj.get("taxonomyId"),
                "matched_seed_entity_pks": uniq_pipe(matched_seed_pks),
                "matched_seed_display_labels": uniq_pipe(metadata.get("seed_display_label") for metadata in matched_seed_metadata),
                "matched_seed_labels": uniq_pipe(entity_label(entity_map.get(pk)) for pk in matched_seed_pks),
                "matched_seed_resolution_status": summarize_resolution_status(matched_seed_statuses),
                "matched_seed_resolution_statuses": uniq_pipe(matched_seed_statuses),
                "subject_seed_resolution_status": seed_metadata_by_pk.get(subject_pk, {}).get("seed_resolution_status", ""),
                "object_seed_resolution_status": seed_metadata_by_pk.get(object_pk, {}).get("seed_resolution_status", ""),
                "other_entity_pks": uniq_pipe(other_pks),
                "other_entity_labels": uniq_pipe(entity_label(entity) for entity in other_entities if entity),
                "other_entity_types": uniq_pipe((entity or {}).get("entityType") for entity in other_entities),
                "other_canonical_identifiers": uniq_pipe((entity or {}).get("canonicalIdentifier") for entity in other_entities),
            }
        )
    return pd.DataFrame(rows).sort_values(["sources", "predicate", "relation_pk"]).reset_index(drop=True)


def fetch_relation_evidence(
    client: OmniPathFastApiClient,
    relation_pks: list[str],
    evidence_scope: str,
    strict: bool = False,
    progress: bool = False,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    evidence_rows: list[dict[str, Any]] = []
    attr_rows: list[dict[str, Any]] = []
    error_rows: list[dict[str, Any]] = []

    total = len(relation_pks)
    for index, relation_pk in enumerate(relation_pks, start=1):
        if progress and (index == 1 or index % 10 == 0 or index == total):
            print(f"{evidence_scope} evidence {index}/{total}", flush=True)
        try:
            payload = client.get(f"/relations/{relation_pk}/evidence")
        except Exception as exc:
            error_rows.append({"evidence_scope": evidence_scope, "relation_pk": relation_pk, "error": str(exc)})
            if strict:
                raise
            continue

        for evidence in payload.get("evidence") or []:
            evidence_rows.append(
                {
                    "evidence_scope": evidence_scope,
                    "relation_pk": relation_pk,
                    "source": evidence.get("source"),
                    "relation_evidence_pk": evidence.get("relationEvidencePk"),
                    "nested_evidence_json": safe_json(evidence.get("evidence")),
                    "n_record_attributes": len(evidence.get("recordAttributes") or []),
                    "n_subject_attributes": len(evidence.get("subjectAttributes") or []),
                    "n_object_attributes": len(evidence.get("objectAttributes") or []),
                }
            )
            for group_name in ("recordAttributes", "subjectAttributes", "objectAttributes"):
                attribute_group = group_name.replace("Attributes", "").lower()
                for attr in evidence.get(group_name) or []:
                    attr_rows.append(
                        {
                            "evidence_scope": evidence_scope,
                            "relation_pk": relation_pk,
                            "source": evidence.get("source"),
                            "relation_evidence_pk": evidence.get("relationEvidencePk"),
                            "attribute_group": attribute_group,
                            "term": attr.get("term"),
                            "value": attr.get("value"),
                            "unit": attr.get("unit"),
                            "scope": attr.get("scope"),
                        }
                    )

    return (
        pd.DataFrame(evidence_rows),
        pd.DataFrame(attr_rows),
        pd.DataFrame(error_rows, columns=["evidence_scope", "relation_pk", "error"]),
    )


def attributes_wide(attributes_long: pd.DataFrame) -> pd.DataFrame:
    if attributes_long.empty:
        return pd.DataFrame(columns=["evidence_scope", "relation_pk", "relation_evidence_pk", "source"])
    grouped = (
        attributes_long.groupby(["evidence_scope", "relation_pk", "relation_evidence_pk", "source", "term"], dropna=False)["value"]
        .apply(uniq_pipe)
        .reset_index()
    )
    wide = grouped.pivot_table(
        index=["evidence_scope", "relation_pk", "relation_evidence_pk", "source"],
        columns="term",
        values="value",
        aggfunc=uniq_pipe,
    ).reset_index()
    wide.columns.name = None
    return wide


def build_hmdb_classes(direct_relations: pd.DataFrame, participant_identifiers: pd.DataFrame) -> pd.DataFrame:
    if direct_relations.empty:
        return pd.DataFrame()
    rows = direct_relations[
        direct_relations["sources"].map(lambda value: source_has(value, "hmdb"))
        & direct_relations["other_canonical_identifiers"].fillna("").str.contains("CHEMONTID", na=False)
    ].copy()
    if rows.empty:
        return pd.DataFrame()

    synonym_map: dict[str, str] = {}
    if not participant_identifiers.empty:
        syn = participant_identifiers[
            participant_identifiers["identifier_type"].fillna("").str.startswith("Synonym")
        ]
        synonym_map = syn.groupby("entity_pk")["identifier"].apply(uniq_pipe).to_dict()

    out = pd.DataFrame(
        {
            "class_label": rows["other_entity_labels"],
            "class_synonym": rows["other_entity_pks"].map(lambda value: synonym_map.get(split_pipe(value)[0], "") if split_pipe(value) else ""),
            "class_id": rows["other_canonical_identifiers"],
            "class_entity_pk": rows["other_entity_pks"],
            "source": rows["sources"],
            "predicate": rows["predicate"],
            "matched_seed_entity_pks": rows["matched_seed_entity_pks"],
            "matched_seed_display_labels": rows["matched_seed_display_labels"],
            "matched_seed_resolution_status": rows["matched_seed_resolution_status"],
            "matched_seed_resolution_statuses": rows["matched_seed_resolution_statuses"],
            "relation_pk": rows["relation_pk"],
        }
    )
    return out.drop_duplicates().sort_values(["class_label", "class_id"]).reset_index(drop=True)


def build_food_occurrence(direct_relations: pd.DataFrame) -> pd.DataFrame:
    if direct_relations.empty:
        return pd.DataFrame()
    mask = direct_relations["sources"].map(lambda value: source_has(value, "foodb")) | direct_relations["other_entity_types"].fillna("").str.contains("Food:OM:0020")
    rows = direct_relations[mask].copy()
    if rows.empty:
        return pd.DataFrame()
    out = pd.DataFrame(
        {
            "food_name": rows["other_entity_labels"],
            "source": rows["sources"],
            "predicate": rows["predicate"],
            "matched_seed_entity_pks": rows["matched_seed_entity_pks"],
            "matched_seed_display_labels": rows["matched_seed_display_labels"],
            "matched_seed_resolution_status": rows["matched_seed_resolution_status"],
            "matched_seed_resolution_statuses": rows["matched_seed_resolution_statuses"],
            "food_entity_pk": rows["other_entity_pks"],
            "food_canonical": rows["other_canonical_identifiers"],
            "relation_pk": rows["relation_pk"],
            "evidence_count": rows["evidence_count"],
        }
    )
    return out.drop_duplicates().sort_values(["food_name", "food_canonical"]).reset_index(drop=True)


def build_metabolic_reactions(direct_relations: pd.DataFrame, direct_attributes_wide: pd.DataFrame) -> pd.DataFrame:
    if direct_relations.empty:
        return pd.DataFrame()
    mask = (
        direct_relations["sources"].map(lambda value: any(source_has(value, src) for src in METABOLIC_SOURCES))
        | direct_relations["other_entity_types"].fillna("").str.contains("Reaction:OM:0015|Transport:OM:0035", regex=True)
    )
    rows = direct_relations[mask].copy()
    if rows.empty:
        return pd.DataFrame()

    attrs = direct_attributes_wide[direct_attributes_wide["evidence_scope"].eq("direct")].copy() if not direct_attributes_wide.empty else pd.DataFrame()
    out = rows.merge(attrs, on="relation_pk", how="left", suffixes=("", "_attr"))
    rename = {
        "other_entity_pks": "linked_entity_pk",
        "other_entity_labels": "linked_entity_label",
        "other_canonical_identifiers": "linked_entity_canonical",
        "Conversion Direction:OM:1211": "conversion_direction",
        "Pathway Participation:OM:0607": "pathway_participation",
        "Reactant:OM:0310": "reactant_flag",
        "Product:OM:0311": "product_flag",
        "Stoichiometry:OM:1226": "stoichiometry",
        "Subcellular Location:OM:0604": "subcellular_location",
    }
    out = out.rename(columns=rename)
    base_cols = [
        "source",
        "sources",
        "predicate",
        "relation_category",
        "matched_seed_entity_pks",
        "matched_seed_display_labels",
        "matched_seed_resolution_status",
        "matched_seed_resolution_statuses",
        "linked_entity_pk",
        "linked_entity_label",
        "linked_entity_canonical",
        "other_entity_types",
        "conversion_direction",
        "pathway_participation",
        "reactant_flag",
        "product_flag",
        "stoichiometry",
        "subcellular_location",
        "relation_pk",
        "relation_evidence_pk",
    ]
    for col in base_cols:
        if col not in out.columns:
            out[col] = pd.NA
    return out[base_cols].drop_duplicates().sort_values(["sources", "predicate", "linked_entity_label", "relation_pk"]).reset_index(drop=True)


def build_macdb_associations(direct_relations: pd.DataFrame, direct_attributes_wide: pd.DataFrame) -> pd.DataFrame:
    if direct_relations.empty:
        return pd.DataFrame()
    rows = direct_relations[direct_relations["sources"].map(lambda value: source_has(value, "macdb"))].copy()
    if rows.empty:
        return pd.DataFrame()

    attrs = direct_attributes_wide[direct_attributes_wide["evidence_scope"].eq("direct")].copy() if not direct_attributes_wide.empty else pd.DataFrame()
    out = rows.merge(attrs, on="relation_pk", how="left", suffixes=("", "_attr"))
    out["subject_name"] = out["other_entity_labels"]
    for raw_col, clean_col in MACDB_TERM_MAP.items():
        out[clean_col] = out[raw_col] if raw_col in out.columns else pd.NA

    cols = [
        "subject_name",
        "matched_seed_entity_pks",
        "matched_seed_display_labels",
        "matched_seed_resolution_status",
        "matched_seed_resolution_statuses",
        "disease_type",
        "disease_subtype",
        "tissue_original",
        "experimental_method",
        "case_sample_count",
        "control_sample_count",
        "pubmed_ids",
        "pmc_ids",
        "study_description",
        "case_description",
        "control_description",
        "contrast_p_value",
        "contrast_logfc",
        "conclusion",
        "metac_id",
        "phenotype_result",
        "relation_pk",
        "relation_evidence_pk",
        "source",
    ]
    for col in cols:
        if col not in out.columns:
            out[col] = pd.NA
    return out[cols].drop_duplicates().sort_values(["disease_type", "subject_name", "tissue_original", "relation_pk"]).reset_index(drop=True)


def normalize_tissue_for_plot(value: Any) -> str | None:
    if pd.isna(value):
        return None
    text = str(value).strip()
    if not text:
        return None
    if "tissue" in text.lower():
        return "Tissue"
    return text


def build_macdb_long(macdb_associations: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    if macdb_associations.empty:
        return pd.DataFrame(), pd.DataFrame()
    rows: list[dict[str, Any]] = []
    exploded_fields = [
        "tissue_original",
        "experimental_method",
        "case_sample_count",
        "control_sample_count",
        "pubmed_ids",
        "pmc_ids",
        "study_description",
        "case_description",
        "control_description",
    ]
    for _, record in macdb_associations.iterrows():
        lists = {field: split_pipe(record.get(field)) for field in exploded_fields}
        n = max([len(values) for values in lists.values()] + [1])
        for values in zip_longest(*[lists[field] for field in exploded_fields], fillvalue=None):
            row = {field: values[i] for i, field in enumerate(exploded_fields)}
            row.update(
                {
                    "subject_name": record.get("subject_name"),
                    "matched_seed_entity_pks": record.get("matched_seed_entity_pks"),
                    "matched_seed_display_labels": record.get("matched_seed_display_labels"),
                    "matched_seed_resolution_status": record.get("matched_seed_resolution_status"),
                    "matched_seed_resolution_statuses": record.get("matched_seed_resolution_statuses"),
                    "disease_type": record.get("disease_type"),
                    "disease_subtype": record.get("disease_subtype"),
                    "relation_pk": record.get("relation_pk"),
                    "relation_evidence_pk": record.get("relation_evidence_pk"),
                    "source": record.get("source"),
                    "contrast_p_value": record.get("contrast_p_value"),
                    "contrast_logfc": record.get("contrast_logfc"),
                    "conclusion": record.get("conclusion"),
                    "metac_id": record.get("metac_id"),
                    "phenotype_result": record.get("phenotype_result"),
                }
            )
            row["tissue_group"] = normalize_tissue_for_plot(row.get("tissue_original"))
            row["study_key"] = "|".join(
                str(row.get(field) or "")
                for field in [
                    "disease_type",
                    "subject_name",
                    "tissue_original",
                    "pubmed_ids",
                    "pmc_ids",
                    "study_description",
                    "relation_evidence_pk",
                ]
            )
            rows.append(row)
            if n == 0:
                break
    long_df = pd.DataFrame(rows).drop_duplicates().reset_index(drop=True)
    cancer_mask = long_df["disease_type"].fillna("").str.contains(
        "cancer|carcinoma|lymphoma|leukemia|leukaemia|melanoma|tumou?r|neoplasm",
        case=False,
        regex=True,
    )
    cancer_long = long_df[cancer_mask & long_df["tissue_group"].notna()].copy()
    if cancer_long.empty:
        return long_df, pd.DataFrame()
    summary = (
        cancer_long.groupby(["disease_type", "tissue_group"], dropna=False)["study_key"]
        .nunique()
        .reset_index(name="evidence_count")
    )
    totals = summary.groupby("disease_type", as_index=False)["evidence_count"].sum().rename(columns={"evidence_count": "total_evidence"})
    summary = summary.merge(totals, on="disease_type").sort_values(["total_evidence", "disease_type", "tissue_group"], ascending=[False, True, True])
    return long_df, summary.reset_index(drop=True)


def identifier_summary(entity_identifiers: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    if entity_identifiers.empty:
        return pd.DataFrame(), pd.DataFrame()
    no_syn = entity_identifiers[
        ~entity_identifiers["identifier_type"].fillna("").str.contains("Synonym", case=False)
    ].copy()
    no_syn["node_category"] = no_syn["identifier_type_clean"].map(identifier_category)
    summary = (
        no_syn.groupby(["identifier_type_clean", "node_category"], dropna=False)
        .agg(
            evidence_count=("identifier", "nunique"),
            n_entities=("entity_pk", "nunique"),
            values=("identifier", lambda values: uniq_pipe(list(OrderedDict((str(v), None) for v in values if pd.notna(v)).keys())[:12])),
        )
        .reset_index()
        .sort_values(["evidence_count", "identifier_type_clean"], ascending=[False, True])
    )
    return no_syn.reset_index(drop=True), summary.reset_index(drop=True)


def build_mindmap_tables(
    entity_identifiers_no_syn: pd.DataFrame,
    hmdb_classes: pd.DataFrame,
    food_occurrence: pd.DataFrame,
    metabolic_reactions: pd.DataFrame,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    center = "Azelaic acid / Azelate"
    rows = [
        {
            "node_id": center,
            "label": center,
            "category": "Center",
            "evidence": 1,
            "source_table": "selected_entities",
            "matched_seed_resolution_status": "resolved",
            "matched_seed_display_labels": "Azelate",
        }
    ]

    if not entity_identifiers_no_syn.empty:
        for _, row in entity_identifiers_no_syn.iterrows():
            label = f"{row['identifier_type_clean']}: {row['identifier']}"
            rows.append(
                {
                    "node_id": f"{row['entity_pk']}::{row['identifier_type_clean']}::{row['identifier']}",
                    "label": label,
                    "category": row["node_category"],
                    "evidence": 1,
                    "source_table": "entity_identifiers",
                    "matched_seed_resolution_status": row.get("seed_resolution_status"),
                    "matched_seed_display_labels": row.get("seed_display_label"),
                }
            )

    if not hmdb_classes.empty:
        for _, row in hmdb_classes.iterrows():
            rows.append(
                {
                    "node_id": str(row.get("class_entity_pk") or row.get("class_id") or row.get("class_label")),
                    "label": row.get("class_label"),
                    "category": "Class",
                    "evidence": 1,
                    "source_table": "hmdb_classes",
                    "matched_seed_resolution_status": row.get("matched_seed_resolution_status"),
                    "matched_seed_display_labels": row.get("matched_seed_display_labels"),
                }
            )

    if not food_occurrence.empty:
        food_summary = (
            food_occurrence.groupby(
                ["food_entity_pk", "food_name", "matched_seed_resolution_status", "matched_seed_display_labels"],
                dropna=False,
            )["relation_pk"]
            .nunique()
            .reset_index(name="evidence")
        )
        for _, row in food_summary.iterrows():
            rows.append(
                {
                    "node_id": str(row.get("food_entity_pk") or row.get("food_name")),
                    "label": row.get("food_name"),
                    "category": "Food",
                    "evidence": int(row.get("evidence") or 1),
                    "source_table": "food_occurrence",
                    "matched_seed_resolution_status": row.get("matched_seed_resolution_status"),
                    "matched_seed_display_labels": row.get("matched_seed_display_labels"),
                }
            )

    if not metabolic_reactions.empty:
        rxn_summary = (
            metabolic_reactions.groupby(
                ["linked_entity_pk", "linked_entity_label", "matched_seed_resolution_status", "matched_seed_display_labels"],
                dropna=False,
            )["relation_pk"]
            .nunique()
            .reset_index(name="evidence")
        )
        for _, row in rxn_summary.iterrows():
            rows.append(
                {
                    "node_id": str(row.get("linked_entity_pk") or row.get("linked_entity_label")),
                    "label": row.get("linked_entity_label"),
                    "category": "Metabolic reaction",
                    "evidence": int(row.get("evidence") or 1),
                    "source_table": "metabolic_reactions",
                    "matched_seed_resolution_status": row.get("matched_seed_resolution_status"),
                    "matched_seed_display_labels": row.get("matched_seed_display_labels"),
                }
            )

    nodes = pd.DataFrame(rows)
    nodes = nodes[nodes["label"].notna() & nodes["label"].astype(str).ne("")]
    nodes = nodes.drop_duplicates(["node_id", "category"]).reset_index(drop=True)
    edges = pd.DataFrame(
        [
            {
                "source": center,
                "target": row["node_id"],
                "target_label": row["label"],
                "category": row["category"],
                "weight": row["evidence"],
                "source_table": row["source_table"],
                "matched_seed_resolution_status": row.get("matched_seed_resolution_status"),
                "matched_seed_display_labels": row.get("matched_seed_display_labels"),
            }
            for _, row in nodes[nodes["category"].ne("Center")].iterrows()
        ]
    )
    return nodes, edges


def build_entity_overview(
    selected_entities: pd.DataFrame,
    relation_query_hits: pd.DataFrame,
    relation_evidence_records: pd.DataFrame,
    relation_attributes_long: pd.DataFrame,
    entity_identifiers: pd.DataFrame,
) -> pd.DataFrame:
    rows = selected_entities.copy()
    relation_counts = (
        relation_query_hits.groupby("query_entity_pk")["relation_pk"].nunique().rename("n_interactions")
        if not relation_query_hits.empty
        else pd.Series(dtype=int)
    )
    identifier_counts = (
        entity_identifiers.groupby("entity_pk")["identifier"].nunique().rename("n_entity_annotations")
        if not entity_identifiers.empty
        else pd.Series(dtype=int)
    )
    structure_counts = (
        entity_identifiers[
            entity_identifiers["identifier_type_clean"].fillna("").str.contains(
                "SMILES|InChI|IUPAC|Formula", case=False, regex=True
            )
        ]
        .groupby("entity_pk")["identifier"]
        .nunique()
        .rename("n_structure_annotations")
        if not entity_identifiers.empty
        else pd.Series(dtype=int)
    )

    if not relation_query_hits.empty and not relation_evidence_records.empty:
        seed_relation_pairs = relation_query_hits[["query_entity_pk", "relation_pk"]].drop_duplicates()
        seed_evidence = seed_relation_pairs.merge(
            relation_evidence_records[["relation_pk", "relation_evidence_pk"]].drop_duplicates(),
            on="relation_pk",
            how="left",
        )
        evidence_counts = seed_evidence.groupby("query_entity_pk")["relation_evidence_pk"].nunique().rename("n_relation_evidence_records")
    else:
        evidence_counts = pd.Series(dtype=int)

    if not relation_query_hits.empty and not relation_attributes_long.empty:
        seed_relation_pairs = relation_query_hits[["query_entity_pk", "relation_pk"]].drop_duplicates()
        seed_attrs = seed_relation_pairs.merge(
            relation_attributes_long[["relation_pk", "relation_evidence_pk", "term", "value"]].drop_duplicates(),
            on="relation_pk",
            how="left",
        )
        attribute_counts = (
            seed_attrs.dropna(subset=["term", "value"])
            .groupby("query_entity_pk")
            .size()
            .rename("n_relation_attribute_annotations")
        )
    else:
        attribute_counts = pd.Series(dtype=int)

    rows["database_status"] = "present in OmniPath metabo DB"
    rows["entity_resolution_status"] = rows["seed_resolution_status"]
    rows["query_entity_label"] = rows["seed_display_label"]
    rows["api_entity_label"] = rows["label"]
    rows["n_interactions"] = rows["entity_pk"].map(relation_counts).fillna(0).astype(int)
    rows["n_relation_evidence_records"] = rows["entity_pk"].map(evidence_counts).fillna(0).astype(int)
    rows["n_relation_attribute_annotations"] = rows["entity_pk"].map(attribute_counts).fillna(0).astype(int)
    rows["n_entity_annotations"] = rows["entity_pk"].map(identifier_counts).fillna(0).astype(int)
    rows["n_structure_annotations"] = rows["entity_pk"].map(structure_counts).fillna(0).astype(int)
    cols = [
        "seed_order",
        "query_entity_label",
        "entity_resolution_status",
        "database_status",
        "entity_pk",
        "api_entity_label",
        "canonical_identifier",
        "canonical_identifier_type",
        "sources",
        "n_interactions",
        "n_relation_evidence_records",
        "n_relation_attribute_annotations",
        "n_entity_annotations",
        "n_structure_annotations",
    ]
    return rows[cols].sort_values("seed_order").reset_index(drop=True)


def build_entity_annotations(entity_identifiers_no_syn: pd.DataFrame, entity_sources: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    if not entity_identifiers_no_syn.empty:
        identifiers = entity_identifiers_no_syn.copy()
        identifiers["annotation_category"] = identifiers["identifier_type_clean"].map(
            lambda value: "Name" if str(value) == "Name" else identifier_category(str(value))
        )
        identifiers["_normalized_identifier"] = identifiers["identifier"].fillna("").astype(str).str.casefold().str.strip()
        identifiers = identifiers.drop_duplicates(["entity_pk", "identifier_type_clean", "_normalized_identifier"])
        for _, row in identifiers.iterrows():
            rows.append(
                {
                    "seed_order": row.get("seed_order"),
                    "query_entity_label": row.get("seed_display_label"),
                    "entity_resolution_status": row.get("seed_resolution_status"),
                    "entity_pk": row.get("entity_pk"),
                    "api_entity_label": row.get("label"),
                    "canonical_identifier": row.get("canonical_identifier"),
                    "annotation_group": "identifier",
                    "annotation_category": row.get("annotation_category"),
                    "annotation_type": row.get("identifier_type_clean"),
                    "annotation_value": row.get("identifier"),
                }
            )
    if not entity_sources.empty:
        for _, row in entity_sources.drop_duplicates(["entity_pk", "source"]).iterrows():
            rows.append(
                {
                    "seed_order": row.get("seed_order"),
                    "query_entity_label": row.get("seed_display_label"),
                    "entity_resolution_status": row.get("seed_resolution_status"),
                    "entity_pk": row.get("entity_pk"),
                    "api_entity_label": row.get("label"),
                    "canonical_identifier": row.get("canonical_identifier"),
                    "annotation_group": "source",
                    "annotation_category": "Source",
                    "annotation_type": "source_membership",
                    "annotation_value": row.get("source"),
                }
            )
    out = pd.DataFrame(rows)
    if out.empty:
        return out
    return out.sort_values(["seed_order", "annotation_group", "annotation_category", "annotation_type", "annotation_value"]).reset_index(drop=True)


def seed_role_for_relation(seed_pk: Any, relation_row: pd.Series) -> str:
    if seed_pk == relation_row.get("subject_entity_pk") and seed_pk == relation_row.get("object_entity_pk"):
        return "subject_and_object"
    if seed_pk == relation_row.get("subject_entity_pk"):
        return "subject"
    if seed_pk == relation_row.get("object_entity_pk"):
        return "object"
    return "matched_seed"


def partner_for_relation(seed_pk: Any, relation_row: pd.Series) -> tuple[str, str, str]:
    if seed_pk == relation_row.get("subject_entity_pk"):
        return (
            relation_row.get("object_label"),
            relation_row.get("object_entity_type"),
            relation_row.get("object_canonical"),
        )
    if seed_pk == relation_row.get("object_entity_pk"):
        return (
            relation_row.get("subject_label"),
            relation_row.get("subject_entity_type"),
            relation_row.get("subject_canonical"),
        )
    return (
        relation_row.get("other_entity_labels"),
        relation_row.get("other_entity_types"),
        relation_row.get("other_canonical_identifiers"),
    )


def build_relations_human(relation_query_hits: pd.DataFrame, direct_relations: pd.DataFrame) -> pd.DataFrame:
    if relation_query_hits.empty or direct_relations.empty:
        return pd.DataFrame()
    merged = relation_query_hits.merge(direct_relations, on="relation_pk", how="left", suffixes=("_query", ""))
    rows: list[dict[str, Any]] = []
    for _, row in merged.iterrows():
        partner_label, partner_type, partner_canonical = partner_for_relation(row.get("query_entity_pk"), row)
        rows.append(
            {
                "query_entity_label": row.get("query_seed_display_label"),
                "entity_resolution_status": row.get("query_seed_resolution_status"),
                "query_entity_pk": row.get("query_entity_pk"),
                "interaction_type": row.get("predicate"),
                "relation_category": row.get("relation_category"),
                "sources": row.get("sources"),
                "source_count": len(split_pipe(row.get("sources"))),
                "seed_role": seed_role_for_relation(row.get("query_entity_pk"), row),
                "partner_label": partner_label,
                "partner_entity_type": partner_type,
                "partner_canonical_identifier": partner_canonical,
                "evidence_count": row.get("evidence_count"),
                "subject_label": row.get("subject_label"),
                "predicate": row.get("predicate"),
                "object_label": row.get("object_label"),
                "relation_pk": row.get("relation_pk"),
            }
        )
    return (
        pd.DataFrame(rows)
        .drop_duplicates()
        .sort_values(["entity_resolution_status", "query_entity_label", "sources", "interaction_type", "partner_label", "relation_pk"])
        .reset_index(drop=True)
    )


def build_interaction_types_by_source(relations: pd.DataFrame) -> pd.DataFrame:
    resolved = relations[relations["entity_resolution_status"].eq("resolved")].copy()
    rows = [
        {
            "interaction_type": row.get("interaction_type"),
            "source": source,
            "relation_pk": row.get("relation_pk"),
            "query_entity_label": row.get("query_entity_label"),
            "entity_resolution_status": row.get("entity_resolution_status"),
        }
        for _, row in resolved.iterrows()
        for source in split_pipe(row.get("sources"))
    ]
    if not rows:
        return pd.DataFrame()
    return (
        pd.DataFrame(rows)
        .groupby(["interaction_type", "source", "query_entity_label", "entity_resolution_status"], dropna=False)["relation_pk"]
        .nunique()
        .reset_index(name="relation_count")
        .sort_values(["relation_count", "interaction_type", "source"], ascending=[False, True, True])
        .reset_index(drop=True)
    )


def cancer_mask(df: pd.DataFrame) -> pd.Series:
    return df["disease_type"].fillna("").str.contains(
        "cancer|carcinoma|lymphoma|leukemia|leukaemia|melanoma|tumou?r|neoplasm",
        case=False,
        regex=True,
    )


def build_cancer_assoc_by_sample_type(macdb_long: pd.DataFrame) -> pd.DataFrame:
    if macdb_long.empty:
        return pd.DataFrame()
    cancer = macdb_long[macdb_long["matched_seed_resolution_status"].eq("resolved") & cancer_mask(macdb_long)].copy()
    if cancer.empty:
        return pd.DataFrame()
    cancer["sample_type"] = cancer["tissue_group"].replace({"A2780 cells": "Cells"})
    study_unit_cols = [
        "disease_type",
        "sample_type",
        "pubmed_ids",
        "pmc_ids",
        "study_description",
        "case_description",
        "control_description",
    ]
    return (
        cancer.drop_duplicates(study_unit_cols)
        .groupby(["disease_type", "sample_type"], dropna=False)
        .size()
        .reset_index(name="evidence_count")
        .sort_values(["evidence_count", "disease_type", "sample_type"], ascending=[False, True, True])
        .reset_index(drop=True)
    )


def build_fig_mindmap_data(mindmap_nodes: pd.DataFrame, mindmap_edges: pd.DataFrame) -> pd.DataFrame:
    if mindmap_nodes.empty:
        return pd.DataFrame()
    nodes = mindmap_nodes[
        mindmap_nodes["category"].eq("Center") | mindmap_nodes["matched_seed_resolution_status"].eq("resolved")
    ].copy()
    keep_node_ids = set(nodes["node_id"])
    edges = mindmap_edges[
        mindmap_edges["source"].isin(keep_node_ids)
        & mindmap_edges["target"].isin(keep_node_ids)
        & mindmap_edges["matched_seed_resolution_status"].eq("resolved")
    ].copy()
    node_rows = [
        {
            "record_type": "node",
            "node_id": row.get("node_id"),
            "label": row.get("label"),
            "category": row.get("category"),
            "evidence": row.get("evidence"),
            "source_node": "",
            "target_node": "",
            "target_label": "",
            "weight": "",
            "source_table": row.get("source_table"),
            "entity_resolution_status": row.get("matched_seed_resolution_status"),
        }
        for _, row in nodes.iterrows()
    ]
    edge_rows = [
        {
            "record_type": "edge",
            "node_id": "",
            "label": "",
            "category": row.get("category"),
            "evidence": "",
            "source_node": row.get("source"),
            "target_node": row.get("target"),
            "target_label": row.get("target_label"),
            "weight": row.get("weight"),
            "source_table": row.get("source_table"),
            "entity_resolution_status": row.get("matched_seed_resolution_status"),
        }
        for _, row in edges.iterrows()
    ]
    return pd.DataFrame(node_rows + edge_rows)


def curate_macdb_associations(macdb_long: pd.DataFrame) -> pd.DataFrame:
    if macdb_long.empty:
        return pd.DataFrame()
    cols = [
        "matched_seed_display_labels",
        "matched_seed_resolution_status",
        "subject_name",
        "disease_type",
        "disease_subtype",
        "tissue_original",
        "tissue_group",
        "experimental_method",
        "case_sample_count",
        "control_sample_count",
        "pubmed_ids",
        "pmc_ids",
        "study_description",
        "case_description",
        "control_description",
        "contrast_p_value",
        "contrast_logfc",
        "conclusion",
        "metac_id",
        "phenotype_result",
        "relation_pk",
        "relation_evidence_pk",
        "source",
    ]
    for col in cols:
        if col not in macdb_long.columns:
            macdb_long[col] = pd.NA
    return macdb_long[cols].drop_duplicates().sort_values(["disease_type", "tissue_group", "subject_name"]).reset_index(drop=True)


def build_supplement_tables(
    selected_entities: pd.DataFrame,
    relation_query_hits: pd.DataFrame,
    relation_evidence_records: pd.DataFrame,
    relation_attributes_long: pd.DataFrame,
    entity_identifiers: pd.DataFrame,
    entity_identifiers_no_syn: pd.DataFrame,
    entity_sources: pd.DataFrame,
    direct_relations: pd.DataFrame,
    hmdb_classes: pd.DataFrame,
    food_occurrence: pd.DataFrame,
    metabolic_reactions: pd.DataFrame,
    macdb_long: pd.DataFrame,
    mindmap_nodes: pd.DataFrame,
    mindmap_edges: pd.DataFrame,
) -> OrderedDict[str, pd.DataFrame]:
    relations = build_relations_human(relation_query_hits, direct_relations)
    tables: OrderedDict[str, pd.DataFrame] = OrderedDict()
    tables["summary"] = build_entity_overview(
        selected_entities,
        relation_query_hits,
        relation_evidence_records,
        relation_attributes_long,
        entity_identifiers,
    )
    tables["entity_annotations"] = build_entity_annotations(entity_identifiers_no_syn, entity_sources)
    tables["relations"] = relations
    tables["hmdb_classes"] = hmdb_classes
    tables["food_occurrence"] = food_occurrence
    tables["metabolic_reactions"] = metabolic_reactions
    tables["macdb_associations"] = curate_macdb_associations(macdb_long)
    tables["interaction_types_by_source"] = build_interaction_types_by_source(relations)
    tables["cancer_assoc_by_sample_type"] = build_cancer_assoc_by_sample_type(macdb_long)
    tables["fig_mindmap_data"] = build_fig_mindmap_data(mindmap_nodes, mindmap_edges)
    return tables


def excel_safe(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    object_cols = out.select_dtypes(include=["object"]).columns
    for col in object_cols:
        out[col] = out[col].map(lambda value: str(value)[:32000] if isinstance(value, str) and len(value) > 32000 else value)
    return out


def save_outputs(tables: dict[str, pd.DataFrame], outdir: Path, figure_paths: dict[str, Path]) -> tuple[Path, Path]:
    csv_dir = outdir / "csvs"
    csv_dir.mkdir(parents=True, exist_ok=True)
    for name, df in tables.items():
        df.to_csv(csv_dir / f"{name}.csv", index=False)

    descriptions = {
        "summary": "Collapsed entity-resolution and count summary. All three queried seed entities are present in the DB; only the InChI-key Azelate entity is treated as resolved.",
        "entity_annotations": "All non-synonym seed-entity annotations used in the supplement: source memberships, external IDs, names, formulas, structures, and other identifiers.",
        "relations": "Human-readable direct relations for each queried seed entity, including seed resolution status, interaction type, source, partner, evidence count, and subject-predicate-object labels.",
        "hmdb_classes": "HMDB/ChemOnt class relations linked to the queried seed entities.",
        "food_occurrence": "FooDB/food occurrence relations linked to the queried seed entities.",
        "metabolic_reactions": "Direct metabolic reaction or transport relations linked to the queried seed entities.",
        "macdb_associations": "MACDB disease-association evidence in long format with original sample/tissue annotations retained.",
        "interaction_types_by_source": "Resolved-only source-by-interaction counts used for use-case figure panel C.",
        "cancer_assoc_by_sample_type": "Resolved-only cancer association study-unit counts by disease type and sample type used for use-case figure panel D.",
        "fig_mindmap_data": "Resolved-only node and edge records used for the azelaic-acid/Azelate mindmap.",
    }
    legend = pd.DataFrame(
        [
            {
                "sheet_name": name[:31],
                "table_name": name,
                "n_rows": len(df),
                "n_columns": df.shape[1],
                "description": descriptions.get(name, ""),
            }
            for name, df in tables.items()
        ]
    )
    if figure_paths:
        for name, path in figure_paths.items():
            legend.loc[len(legend)] = {
                "sheet_name": "",
                "table_name": f"figure:{name}",
                "n_rows": "",
                "n_columns": "",
                "description": str(path),
            }

    workbook_path = outdir / "azelaic_acid_tables.xlsx"
    used_sheet_names: set[str] = {"legend"}
    with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
        legend.to_excel(writer, sheet_name="legend", index=False)
        for name, df in tables.items():
            base = name[:31]
            sheet_name = base
            suffix = 1
            while sheet_name in used_sheet_names:
                suffix_text = f"_{suffix}"
                sheet_name = f"{base[:31-len(suffix_text)]}{suffix_text}"
                suffix += 1
            used_sheet_names.add(sheet_name)
            excel_safe(df).to_excel(writer, sheet_name=sheet_name, index=False)
        style_workbook(writer.book)
    return csv_dir, workbook_path


def style_workbook(workbook: Any) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        if worksheet.max_row and worksheet.max_column:
            worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for column_cells in worksheet.columns:
            max_len = 0
            for cell in column_cells[:200]:
                value = "" if cell.value is None else str(cell.value)
                max_len = max(max_len, min(len(value), 80))
            width = min(max(max_len + 2, 12), 55)
            worksheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width
        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)


def make_manifest(
    outdir: Path,
    workbook_path: Path,
    csv_dir: Path,
    figure_paths: dict[str, Path],
    tables: dict[str, pd.DataFrame],
    api_base: str,
    selected_entities: pd.DataFrame,
) -> Path:
    manifest = {
        "api_base": api_base,
        "run_utc": datetime.now(timezone.utc).isoformat(),
        "output_dir": str(outdir),
        "csv_dir": str(csv_dir),
        "workbook_path": str(workbook_path),
        "figures": {name: str(path) for name, path in figure_paths.items()},
        "selected_entity_pks": selected_entities["entity_pk"].tolist() if not selected_entities.empty else [],
        "table_counts": {name: {"rows": int(len(df)), "columns": int(df.shape[1])} for name, df in tables.items()},
    }
    manifest_path = outdir / "manifest.json"
    manifest_path.write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    return manifest_path


# Paths relative to the repository root (this file: analyses/azelate/).
REPO_ROOT = Path(__file__).resolve().parents[2]
OUTDIR = REPO_ROOT / "data" / "derived" / "azelate"


def run_pipeline(
    api_base: str = API_BASE,
    outdir: str | Path = OUTDIR,
    query_terms: list[str] | None = None,
    progress: bool = False,
) -> PipelineResult:
    outdir = Path(outdir)
    outdir.mkdir(parents=True, exist_ok=True)

    client = OmniPathFastApiClient(api_base)
    query_terms = query_terms or QUERY_TERMS

    if progress:
        print("Searching and resolving seed entities", flush=True)
    entity_search_hits, search_entity_map = search_entities(client, query_terms)
    entity_resolve_matches, entity_resolve_candidates, resolve_entity_map = resolve_entities(client, query_terms)
    selected_entity_map = select_seed_entities(entity_search_hits, search_entity_map, resolve_entity_map)
    selected_entity_pks = list(selected_entity_map)
    selected_entity_map.update(hydrate_entities(client, selected_entity_pks))
    seed_metadata_by_pk = {pk: seed_selection_metadata(entity) for pk, entity in selected_entity_map.items()}

    seed_order = {str(rule["seed_key"]): idx for idx, rule in enumerate(SEED_SELECTION_RULES, start=1)}
    selected_rows = []
    for entity in selected_entity_map.values():
        metadata = seed_selection_metadata(entity)
        selected_rows.append(
            {
                **metadata,
                "seed_order": seed_order.get(metadata.get("seed_key"), 999),
                **entity_summary(entity),
            }
        )
    selected_entities = pd.DataFrame(selected_rows)
    selected_entities = selected_entities.sort_values(["seed_order", "entity_pk"]).reset_index(drop=True)
    entity_identifiers = add_seed_metadata_columns(
        flatten_identifiers(list(selected_entity_map.values()), "selected_seed"),
        seed_metadata_by_pk,
    )
    entity_sources = add_seed_metadata_columns(
        explode_entity_sources(list(selected_entity_map.values()), "selected_seed"),
        seed_metadata_by_pk,
    )
    entity_identifiers_no_syn, _ = identifier_summary(entity_identifiers)

    if progress:
        print(f"Selected {len(selected_entity_pks)} seed entities; fetching direct relations", flush=True)
    relation_query_hits, direct_relation_map = fetch_relations_for_entities(client, selected_entity_pks)
    relation_query_hits = add_query_seed_metadata(relation_query_hits, seed_metadata_by_pk)
    direct_participant_pks = relation_participant_pks(direct_relation_map)
    participant_entity_map = hydrate_entities(client, direct_participant_pks)
    participant_entity_map.update(selected_entity_map)
    direct_relations = build_direct_relations(direct_relation_map, set(selected_entity_pks), seed_metadata_by_pk, participant_entity_map)

    if progress:
        print(f"Fetched {len(direct_relation_map)} direct relations; fetching direct relation evidence", flush=True)
    participant_identifiers = flatten_identifiers(list(participant_entity_map.values()), "direct_relation_participant")

    relation_evidence_records, relation_attributes_long, _ = fetch_relation_evidence(
        client,
        direct_relations["relation_pk"].dropna().tolist() if not direct_relations.empty else [],
        "direct",
        progress=progress,
    )
    relation_attributes_wide = attributes_wide(relation_attributes_long)
    relation_evidence_records = add_relation_seed_metadata(relation_evidence_records, direct_relations)
    relation_attributes_long = add_relation_seed_metadata(relation_attributes_long, direct_relations)
    relation_attributes_wide = add_relation_seed_metadata(relation_attributes_wide, direct_relations)

    if progress:
        print("Deriving curated supplement tables", flush=True)
    hmdb_classes = build_hmdb_classes(direct_relations, participant_identifiers)
    food_occurrence = build_food_occurrence(direct_relations)
    metabolic_reactions = build_metabolic_reactions(direct_relations, relation_attributes_wide)
    macdb_associations = build_macdb_associations(direct_relations, relation_attributes_wide)
    macdb_long, _ = build_macdb_long(macdb_associations)
    mindmap_nodes, mindmap_edges = build_mindmap_tables(entity_identifiers_no_syn, hmdb_classes, food_occurrence, metabolic_reactions)
    tables = build_supplement_tables(
        selected_entities=selected_entities,
        relation_query_hits=relation_query_hits,
        relation_evidence_records=relation_evidence_records,
        relation_attributes_long=relation_attributes_long,
        entity_identifiers=entity_identifiers,
        entity_identifiers_no_syn=entity_identifiers_no_syn,
        entity_sources=entity_sources,
        direct_relations=direct_relations,
        hmdb_classes=hmdb_classes,
        food_occurrence=food_occurrence,
        metabolic_reactions=metabolic_reactions,
        macdb_long=macdb_long,
        mindmap_nodes=mindmap_nodes,
        mindmap_edges=mindmap_edges,
    )

    figure_paths: dict[str, Path] = {}

    if progress:
        print("Saving CSV and Excel outputs", flush=True)
    csv_dir, workbook_path = save_outputs(tables, outdir, figure_paths)
    manifest_path = make_manifest(outdir, workbook_path, csv_dir, figure_paths, tables, api_base, selected_entities)

    return PipelineResult(
        tables=tables,
        workbook_path=workbook_path,
        csv_dir=csv_dir,
        figure_paths=figure_paths,
        manifest_path=manifest_path,
        output_dir=outdir,
    )


if __name__ == "__main__":
    result = run_pipeline(progress=True)
    print(f"Workbook: {result.workbook_path}")
    print(f"CSV directory: {result.csv_dir}")
    print(f"Manifest: {result.manifest_path}")
    for name, path in result.figure_paths.items():
        print(f"{name}: {path}")
